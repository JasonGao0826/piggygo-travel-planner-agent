"""把报告内容写成 Word（设计规格 17.3 节）：原生可编辑表格、按地点配色、横向 A4、中文字体。"""

import dataclasses
import re
from pathlib import Path
from urllib.parse import urlparse

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from .model import Bullets, Card, Heading, Paragraph, Report, Table

FONT = "Microsoft YaHei"
ACCENT = "2566A8"   # 标题、表头：品牌天蓝加深一档（和网页的强调色一致）
COCOA = RGBColor(0x6B, 0x33, 0x24)   # logo 里"Piggy"的巧克力棕
LOGO = Path(__file__).parent / "assets" / "logo.png"
BRAND = "小猪游记 PiggyGo"
SLOGAN = "交给猪猪，你负责开心出发！"
MUTED = RGBColor(0x59, 0x59, 0x59)
LINK = RGBColor(0x05, 0x63, 0xC1)
LABEL_FILL = "F2F2F2"


def _font(element, size: float | None = None, bold: bool | None = None, color: RGBColor | None = None) -> None:
    """设置字体。中文要单独设置 eastAsia，否则 Word 会用默认的宋体。"""
    font = element.font
    font.name = FONT
    rpr = element.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    if size is not None:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if color is not None:
        font.color.rgb = color


def _shade(element, fill: str) -> None:
    """给单元格或段落加底色。"""
    props = element.get_or_add_tcPr() if element.tag == qn("w:tc") else element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


URL = re.compile(r"https?://[^\s｜，。、）)]+")


def link_label(url: str) -> str:
    """长链接在报告里只显示网站名，例如 airbnb.com。"""
    return urlparse(url).netloc.removeprefix("www.") or url


def _write(paragraph, text: str, size: float = 9, bold: bool = False, color: RGBColor | None = None) -> None:
    """写入文本，"\\n" 变成换行，网址变成可以点的链接。"""
    lines = text.split("\n")
    for n, line in enumerate(lines):
        pos = 0
        for m in URL.finditer(line):
            if m.start() > pos:
                _font(paragraph.add_run(line[pos:m.start()]), size, bold, color)
            _hyperlink(paragraph, m.group(), link_label(m.group()), size)
            pos = m.end()
        run = paragraph.add_run(line[pos:])
        _font(run, size, bold, color)
        if n < len(lines) - 1:
            run.add_break(WD_BREAK.LINE)


def _hyperlink(paragraph, url: str, text: str, size: float = 9) -> None:
    rel_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        fonts.set(qn(attr), FONT)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), str(LINK))
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.extend([fonts, color, underline, sz])
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _repeat_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def _no_split(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))


def _page_number(footer) -> None:
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _write(p, f"{BRAND} · {SLOGAN}    第 ", size=8, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    run.append(t)
    field.append(run)
    p._p.append(field)
    _write(p, " 页", size=8, color=MUTED)


class DocxRenderer:
    def __init__(self) -> None:
        self.doc = Document()
        section = self.doc.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Mm(297), Mm(210)
        for side in ("left_margin", "right_margin"):
            setattr(section, side, Mm(15))
        section.top_margin = section.bottom_margin = Mm(14)
        _page_number(section.footer)

        styles = self.doc.styles
        _font(styles["Normal"], 10)
        styles["Normal"].paragraph_format.space_after = Pt(3)
        for name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
            _font(styles[name], size, True, RGBColor.from_string(ACCENT))
        styles["Heading 1"].paragraph_format.space_before = Pt(12)
        for name in ("Heading 1", "Heading 2", "Heading 3"):
            styles[name].paragraph_format.keep_with_next = True
        styles["List Bullet"].paragraph_format.space_after = Pt(1)

    # ---------- 各类内容块 ----------

    def paragraph(self, block: Paragraph) -> None:
        p = self.doc.add_paragraph()
        _write(p, block.text, size=9 if block.muted else 10, color=MUTED if block.muted else None)

    def bullets(self, block: Bullets) -> None:
        for item in block.items:
            p = self.doc.add_paragraph(style="List Bullet")
            _write(p, item, size=10)

    def heading(self, block: Heading) -> None:
        p = self.doc.add_heading(level=2)
        _write(p, block.text, size=12, bold=True, color=RGBColor.from_string(ACCENT))
        if block.fill:
            _shade(p._p, block.fill)

    @staticmethod
    def _keep_together(table) -> None:
        """表格尽量不跨页：除最后一行外，每行都和下一行保持在同一页（放不下一页的长表格仍会正常分页）。"""
        for row in table.rows[:-1]:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True

    def table(self, block: Table) -> None:
        table = self.doc.add_table(rows=1, cols=len(block.columns))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        header = table.rows[0]
        _repeat_header(header)
        for cell, title in zip(header.cells, block.columns):
            _shade(cell._tc, ACCENT)
            _write(cell.paragraphs[0], title, size=9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        for row in block.rows:
            cells = table.add_row()
            _no_split(cells)
            for cell, text in zip(cells.cells, row.cells):
                if row.fill:
                    _shade(cell._tc, row.fill)
                _write(cell.paragraphs[0], text, size=9, bold=row.bold)
        widths = block.widths_mm or [267 / len(block.columns)] * len(block.columns)
        for row in table.rows:
            for cell, w in zip(row.cells, widths):
                cell.width = Mm(w)
        if len(block.rows) <= 18:
            self._keep_together(table)
        self.doc.add_paragraph()

    def card(self, block: Card) -> None:
        p = self.doc.add_heading(level=3)
        _write(p, block.title, size=11, bold=True, color=RGBColor.from_string(ACCENT))
        if block.intro:
            intro = self.doc.add_paragraph()
            intro.paragraph_format.keep_with_next = True  # 简介和下面的表格放在同一页
            _write(intro, block.intro, size=9, color=MUTED)
        rows = len(block.facts) + (1 if block.links else 0)
        table = self.doc.add_table(rows=rows, cols=2)
        table.style = "Table Grid"
        table.autofit = False
        for row, (label, value) in zip(table.rows, block.facts):
            _no_split(row)
            _shade(row.cells[0]._tc, block.fill or LABEL_FILL)
            _write(row.cells[0].paragraphs[0], label, size=9, bold=True)
            _write(row.cells[1].paragraphs[0], value, size=9)
        if block.links:
            row = table.rows[-1]
            _shade(row.cells[0]._tc, block.fill or LABEL_FILL)
            _write(row.cells[0].paragraphs[0], "链接", size=9, bold=True)
            p = row.cells[1].paragraphs[0]
            for n, link in enumerate(block.links):
                if n:
                    _write(p, "  ·  ", size=9, color=MUTED)
                _hyperlink(p, link.url, link.label)
        for row in table.rows:
            row.cells[0].width, row.cells[1].width = Mm(28), Mm(239)
        self._keep_together(table)
        self.doc.add_paragraph()

    # ---------- 整份文档 ----------

    def cover(self, report: Report) -> None:
        if LOGO.exists():
            logo = self.doc.add_paragraph()
            logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            logo.add_run().add_picture(str(LOGO), height=Mm(48))
        else:
            for _ in range(4):
                self.doc.add_paragraph()
        title = self.doc.add_paragraph(style="Title")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _write(title, report.title, size=24, bold=True, color=RGBColor.from_string(ACCENT))
        for line in report.cover_lines:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _write(p, line, size=12)
        slogan = self.doc.add_paragraph()
        slogan.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _write(slogan, f"{BRAND}｜{SLOGAN}", size=11, bold=True, color=COCOA)
        if report.legend:
            self.doc.add_paragraph()
            label = self.doc.add_paragraph()
            label.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _write(label, "地点配色", size=9, color=MUTED)
            legend = self.doc.add_table(rows=1, cols=len(report.legend))
            legend.alignment = WD_TABLE_ALIGNMENT.CENTER
            for cell, (city, fill) in zip(legend.rows[0].cells, report.legend):
                _shade(cell._tc, fill)
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                _write(cell.paragraphs[0], city, size=9)
                cell.width = Mm(28)

        self.doc.add_page_break()
        toc = self.doc.add_paragraph()
        _write(toc, "目录", size=16, bold=True, color=RGBColor.from_string(ACCENT))
        for s in report.sections:
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Mm(4)
            _write(p, s.title, size=11)
        self.doc.add_page_break()

    def render(self, report: Report, path: str | Path) -> Path:
        self.cover(report)
        for n, section in enumerate(report.sections):
            if section.new_page and n > 0:
                self.doc.add_page_break()
            h = self.doc.add_heading(level=1)
            _write(h, section.title, size=16, bold=True, color=RGBColor.from_string(ACCENT))
            for block in section.blocks:
                {Paragraph: self.paragraph, Bullets: self.bullets, Heading: self.heading,
                 Table: self.table, Card: self.card}[type(block)](block)
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(path)
        return path


_XML_ILLEGAL = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\x7f\ufffe\uffff]")


def _clean(value):
    """Word（XML）不接受控制字符：报告里所有文字先去掉它们。"""
    if isinstance(value, str):
        return _XML_ILLEGAL.sub("", value)
    if isinstance(value, list):
        return [_clean(v) for v in value]
    if isinstance(value, tuple):
        return tuple(_clean(v) for v in value)
    if isinstance(value, dict):
        return {k: _clean(v) for k, v in value.items()}
    if dataclasses.is_dataclass(value) and not isinstance(value, type):
        return dataclasses.replace(value, **{f.name: _clean(getattr(value, f.name))
                                             for f in dataclasses.fields(value) if f.init})
    return value


def render_docx(report: Report, path: str | Path) -> Path:
    return DocxRenderer().render(_clean(report), path)
