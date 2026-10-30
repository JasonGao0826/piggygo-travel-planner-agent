"""Word 报告：内容取自行程数据；同一个时间在所有章节里一致；能生成可编辑的 Word 文件。"""

from datetime import date

from docx import Document

from tripagent.report import build_report, render_docx
from tripagent.report.model import Card, Table
from tripagent.service import TripService

from .conftest import rebuild, nz
from .test_service import TRIP, _v3_milford

TODAY = date(2026, 10, 22)


def section(report, name):
    return next(s for s in report.sections if name in s.title)


def test_sections_follow_v3_report_structure(trip, places):
    titles = [s.title for s in build_report(trip, places, generated_on=TODAY).sections]
    assert titles == [
        "1. 我做了哪些假设 / 待确认事项", "2. 机票", "3. 签证与入境", "4. 钱和支付", "5. 租车与驾照",
        "6. 住宿", "7. 全部游玩项目", "8. 项目详情", "9. 10 天行程总览", "10. 每天具体时间安排", "11. 还没订的项目",
        "12. 待办事项总表", "13. 出发前清单", "14. 已完成的预订",
    ]


def test_all_activities_table_like_v4(trip, places):
    """案例 V4 的"全部游玩项目汇总表"：签到时间、游玩时段和时长、集合点、状态。"""
    cruise = trip.find_item("d3_cruise")
    stop = trip.find_item("d3_mirror_lakes")
    stop.optional, stop.optional_reason = True, "时间不够就跳过"
    table = next(b for b in section(build_report(rebuild(trip), places), "全部游玩项目").blocks if isinstance(b, Table))
    rows = {r.cells[1]: r for r in table.rows}
    row = rows[cruise.title]
    assert row.cells[2].endswith("签到") and "（1 小时 40 分）" in row.cells[3] and row.cells[5].startswith("已订")
    assert row.bold
    assert rows[stop.title].cells[5].startswith("备选\n时间不够就跳过")
    schedule = section(build_report(rebuild(trip), places), "每天具体时间安排")
    assert any(r.cells[1] == stop.title and r.cells[4] == "备选" for b in schedule.blocks if isinstance(b, Table) for r in b.rows)


def test_car_contract_details(trip, places):
    car = next(v for v in trip.vehicles if v.id == "car_south")
    car.insurance, car.dropoff_desk_hours = "全险（零自付额）", "07:00-21:00"
    table = [b for b in section(build_report(rebuild(trip), places), "租车与驾照").blocks if isinstance(b, Table)][2]
    assert table.columns[:3] == ["用车", "保险", "自付额"]
    row = table.rows[0].cells
    assert row[1] == "全险（零自付额）" and row[2] == "待确认" and row[5] == "取 待确认\n还 07:00-21:00"


def test_cover_route_is_derived_from_flights_and_segments(trip, places):
    report = build_report(trip, places, generated_on=TODAY)
    assert "整体路线：墨尔本 ✈ 皇后镇 → 南岛自驾 → 基督城 ✈ 汉密尔顿 → 北岛自驾 → 奥克兰 ✈ 墨尔本" in report.cover_lines
    assert "时间：2026 年 11 月 23 日出发，12 月 2 日返回，共 10 天 9 晚" in report.cover_lines


def test_overview_has_one_row_per_day_with_drive_hours(trip, places):
    table = next(b for b in section(build_report(trip, places), "行程总览").blocks if isinstance(b, Table))
    assert len(table.rows) == 10
    milford = table.rows[2].cells
    assert milford[0] == "D3 11/25 周三" and milford[3] == "约 9.2 小时"
    assert "13:15–14:55 RealNZ 米尔福德峡湾游船（已订）" in milford[2]


def test_card_uses_knowledge_base_facts(trip, places):
    cards = [b for b in section(build_report(trip, places), "项目详情").blocks if isinstance(b, Card)]
    waiotapu = next(c for c in cards if "Wai-O-Tapu" in c.title)
    facts = dict(waiotapu.facts)
    assert waiotapu.title.endswith("（需预约）")
    assert facts["开放时间"] == "08:30–16:30，最晚 15:00 入场"
    assert facts["参考价格"] == "成人 NZ$47/人"


def test_split_groups_show_who_is_going(trip, places):
    cards = [b for b in section(build_report(trip, places), "项目详情").blocks if isinstance(b, Card)]
    hydro = dict(next(c for c in cards if "Hydro Attack" in c.title).facts)
    assert hydro["本次安排"] == "11/24 周二 12:10–12:35（Tony、天哥） 已订\n11/24 周二 13:50–14:15（Nina、Leo） 已订"


def test_report_shows_new_cruise_time_everywhere_after_voucher(trip, decisions, places, drive_times, store):
    """回放第 10 步：凭证写的是 13:15。写入新版本后，报告里任何地方都不能再出现 13:55。"""
    service = TripService(store, places, drive_times)
    service.create_trip(_v3_milford(trip), decisions)
    before = build_report(service.current(TRIP), places).plain_text()
    assert "13:55" in before

    patch = service.propose(TRIP, [{"op": "update", "collection": "items", "id": "d3_cruise", "set": {
        "start": nz("2026-11-25", "13:15").isoformat(), "end": nz("2026-11-25", "14:55").isoformat(),
        "booking_status": "booked", "locked": True, "lock_reason": "evidence", "source": "evidence"}}],
        summary="按游船凭证改为 13:15 场")
    version = service.confirm(patch.id, confirmed_by="m_nina")

    text = build_report(version.state, places, changelog=version.changelog).plain_text()
    # 变更说明里允许出现旧时间，其余提到这次游船的地方都只能是 13:15
    cruise_lines = [line for line in text.splitlines()
                    if "RealNZ 米尔福德峡湾游船" in line and not line.startswith(("修改", "带上："))]
    assert len(cruise_lines) >= 3  # 总览、每日时间表、已订清单
    assert all("13:15–14:55" in line and "13:55" not in line for line in cruise_lines)
    assert "计划 V2" in text and "开始时间 11/25 13:55 → 11/25 13:15" in text


def test_docx_is_editable_tables(trip, places, tmp_path):
    path = render_docx(build_report(trip, places, generated_on=TODAY), tmp_path / "report.docx")
    doc = Document(path)
    assert len(doc.tables) > 30  # 原生表格，不是图片
    assert len(doc.inline_shapes) == 1  # 只有封面的 logo；表格都是原生表格，不是图片
    texts = {cell.text for t in doc.tables for row in t.rows for cell in row.cells}
    assert "Lighthorse Easy Rider 骑马" in texts
    assert doc.sections[0].page_width > doc.sections[0].page_height  # 横向 A4
